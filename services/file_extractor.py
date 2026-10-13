"""Local text extraction for supported job-description and CV files."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}
LEGACY_EXTENSIONS = {".doc"}
MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024


class ExtractionError(ValueError):
    """A user-correctable problem while validating or extracting a file."""


@dataclass(frozen=True)
class ExtractionResult:
    path: Path
    file_type: str
    text: str
    warning: str | None = None


def extract_text_content(text: str, source_name: str = "pasted text") -> ExtractionResult:
    """Validate user-provided text such as a pasted job description."""

    normalized_text = text.strip()
    if not normalized_text:
        raise ExtractionError(f"The {source_name} is empty.")
    return ExtractionResult(path=Path(source_name), file_type="TEXT", text=normalized_text)


def _validate_path(path: Path) -> None:
    if not path.exists():
        raise ExtractionError("The selected file no longer exists.")
    if not path.is_file():
        raise ExtractionError("Please select a file, not a folder.")
    if path.stat().st_size == 0:
        raise ExtractionError("The selected file is empty.")
    if path.stat().st_size > MAX_FILE_SIZE_BYTES:
        raise ExtractionError("The selected file is larger than the 10 MB limit.")


def _extract_pdf(path: Path) -> tuple[str, str | None]:
    try:
        import pymupdf
    except ImportError as exc:
        raise ExtractionError("PDF support is unavailable: install PyMuPDF.") from exc

    try:
        with pymupdf.open(path) as document:
            text = "\n\n".join(page.get_text("text") for page in document).strip()
    except Exception as exc:  # parser-specific exceptions vary by PyMuPDF version
        raise ExtractionError(f"Could not read the PDF: {exc}") from exc

    if not text:
        return "", "This PDF has no selectable text. OCR or a text-based PDF/TXT is needed."
    return text, None


def _extract_docx(path: Path) -> tuple[str, str | None]:
    try:
        from docx import Document
    except ImportError as exc:
        raise ExtractionError("DOCX support is unavailable: install python-docx.") from exc

    try:
        document = Document(path)
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        tables = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    except Exception as exc:
        raise ExtractionError(f"Could not read the DOCX file: {exc}") from exc

    text = "\n".join(paragraphs + tables).strip()
    if not text:
        raise ExtractionError("The DOCX file contains no extractable text.")
    return text, None


def extract_text(file_path: str | Path) -> ExtractionResult:
    """Validate and extract text without modifying or uploading the source file."""

    path = Path(file_path)
    _validate_path(path)
    extension = path.suffix.lower()

    if extension in LEGACY_EXTENSIONS:
        raise ExtractionError(
            "Legacy .doc files are not read directly. Convert the file to DOCX or TXT, "
            "or install a separate LibreOffice conversion step."
        )
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS | LEGACY_EXTENSIONS))
        raise ExtractionError(f"Unsupported file type '{extension or 'none'}'. Supported types: {supported}.")

    if extension == ".pdf":
        text, warning = _extract_pdf(path)
        file_type = "PDF"
    elif extension == ".docx":
        text, warning = _extract_docx(path)
        file_type = "DOCX"
    else:
        try:
            text = path.read_text(encoding="utf-8-sig").strip()
        except UnicodeDecodeError as exc:
            raise ExtractionError("The TXT file is not valid UTF-8 text.") from exc
        if not text:
            raise ExtractionError("The TXT file contains no text.")
        warning = None
        file_type = "TXT"

    return ExtractionResult(path=path, file_type=file_type, text=text, warning=warning)
